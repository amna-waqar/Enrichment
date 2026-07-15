"""Final workbook: merge ALL enrichment sources, tag each match with its source,
re-bucket on exact-domain matches, and refresh before/after coverage.

Sources (priority order; richest first):
  1. FullEnrich  (Chargebee_Company_Enrichment) — full company firmographics
  2. Other_Paying / High_LTV / Activated_Freemiums — past person-level exports;
     company domain comes from the FullEnrich email column (dedicated domain
     columns are ~96% empty), company data is mostly name-only.

Matching is identical to before: exact normalized domain (customer domain or
non-free email domain) == an enrichment domain => 100% confidence. FullEnrich
name-only matches are kept as "review" (never re-bucketed). No fuzzy guesses.
"""

from __future__ import annotations

import argparse
import re

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import build_icp_sheet as base
import match_enrichment as me
import rebucket_enrichment as rb

# enrichment columns, with enr_source added
FINAL_ENRICH_COLS = ["enr_status", "enr_match_key", "enr_confidence", "enr_source",
                     "enr_matched_domain"] + [f"enr_{f}" for f in me.ENRICH_FIELDS]
META_COLS = ["original_bucket", "bucket_source", "moved"]

# LinkedIn-schema column mapping for the 3 past-enrichment files
LINKEDIN_MAP = {
    "Company Name": "Company Name (Linkedin)",
    "Company Linkedin URL": "Company LinkedIn Url",
    "Company Description": "Company Description (Linkedin)",
    "Year Founded": "Company Year Founded (Linkedin)",
    "Company headcount": "Company Headcount (Linkedin)",
    "Company headcount range": "Company Headcount Range (Linkedin)",
    "Company type": "Company Type (Linkedin)",
    "Company industry": "Company Industry (Linkedin)",
}
LINKEDIN_DOMAIN_COLS = ["Domain (FullEnrich)", "Company Website (Linkedin)", "Website"]


def _decode_band(v):
    s = "" if v is None else str(v).strip()
    if s in ("10001", "10000+"):
        return "10001+"
    return rb.decode_band(v)


def _clean(v):
    return "" if (v is None or (isinstance(v, float) and v != v) or str(v).strip().lower() == "nan") else str(v).strip()


def _email_domain(v):
    v = str(v)
    return me.norm_domain(v.split("@")[-1]) if "@" in v else ""


def _completeness(rec: dict) -> int:
    return sum(1 for f in me.ENRICH_FIELDS if rec.get(f, "").strip())


def build_multi_index(fullenrich_path: str, extra_files: list[tuple[str, str]]):
    """Return dom_best, name_unique, dom_band, dom_source, stats."""
    # --- 1. FullEnrich (primary) ---
    dom_best, name_unique, _ = me.build_enrichment_index(fullenrich_path)
    # decode FullEnrich band + record source
    xl = pd.ExcelFile(fullenrich_path)
    fe = pd.concat([pd.read_excel(fullenrich_path, sheet_name=s, dtype=str) for s in xl.sheet_names],
                   ignore_index=True)
    fe["_dom"] = fe["Company Domain"].map(me.norm_domain)
    fe["_band"] = fe["Company headcount range"].map(_decode_band)
    dom_band, dom_source = {}, {}
    band_by_dom = {}
    for _, r in fe.iterrows():
        d = r["_dom"]
        if d and r["_band"] and d not in band_by_dom:
            band_by_dom[d] = r["_band"]
    for d in dom_best:
        dom_band[d] = band_by_dom.get(d)
        dom_source[d] = "FullEnrich"
        dom_best[d]["Company headcount range"] = band_by_dom.get(d) or ""

    # --- 2. extra files (secondary) ---
    for label, path in extra_files:
        xl = pd.ExcelFile(path)
        d = pd.concat([pd.read_excel(path, sheet_name=s, dtype=str) for s in xl.sheet_names],
                      ignore_index=True)
        for _, row in d.iterrows():
            # candidate company domains for this row: explicit column(s) + email domain
            cands = []
            for c in LINKEDIN_DOMAIN_COLS:
                if c in d.columns:
                    nd = me.norm_domain(row.get(c))
                    if nd:
                        cands.append(nd)
            if "Email (FullEnrich)" in d.columns:
                ed = _email_domain(row.get("Email (FullEnrich)"))
                if ed:
                    cands.append(ed)
            cands = [x for x in dict.fromkeys(cands) if x and x not in me.FREE_DOMAINS]
            if not cands:
                continue
            band = _decode_band(row.get("Company Headcount Range (Linkedin)"))
            rec = {f: "" for f in me.ENRICH_FIELDS}
            for tgt, src in LINKEDIN_MAP.items():
                rec[tgt] = _clean(row.get(src))
            rec["Company headcount range"] = band or ""
            hq = ", ".join(x for x in [
                _clean(row.get("Company Headquarters City (Linkedin)")),
                _clean(row.get("Company Headquarters Region (Linkedin)")),
                _clean(row.get("Company Headquarters Country (Linkedin)"))] if x)
            rec["Company headquarters location"] = hq
            for dom in cands:
                r = dict(rec); r["Company Domain"] = dom
                if dom not in dom_best:
                    dom_best[dom] = r
                    dom_source[dom] = label
                    dom_band[dom] = band
                elif dom_source.get(dom) == label and _completeness(r) > _completeness(dom_best[dom]):
                    dom_best[dom] = r
                    dom_band[dom] = band or dom_band.get(dom)

    stats = {"total_domains": len(dom_best)}
    return dom_best, name_unique, dom_band, dom_source, stats


def enrich(df, dom_best, name_unique, dom_band, dom_source):
    recs = []
    for _, row in df.iterrows():
        out = me.match_customer(row, dom_best, name_unique)
        md = out.get("enr_matched_domain") or ""
        out["enr_source"] = dom_source.get(md, "") if out["enr_status"] == "Enriched" else ""
        orig = row["size_bucket"]
        band = dom_band.get(md) if md else None
        if out.get("enr_confidence") != "100%":
            band = None
        if band in rb.NUM_BUCKETS:
            final = band
            source = "FullEnrich/enrichment" if final != orig else "Chargebee (confirmed)"
        else:
            final = orig
            source = "Chargebee (original)"
        out["original_bucket"] = orig
        out["size_bucket_final"] = final
        out["bucket_source"] = source
        out["moved"] = "Yes" if final != orig else ""
        recs.append(out)
    add = pd.DataFrame(recs)
    return pd.concat([df.reset_index(drop=True), add.reset_index(drop=True)], axis=1)


def coverage_by_source(en):
    e = en[en["enr_status"] == "Enriched"]
    out = (e.groupby("enr_source").size().reset_index(name="Accounts enriched")
           .sort_values("Accounts enriched", ascending=False).reset_index(drop=True))
    out.loc[len(out)] = ["TOTAL", len(e)]
    return out


def overview(en, moved, size_b, size_a, dom_source):
    n = len(en)
    ex = int((en["enr_confidence"] == "100%").sum())
    rv = int((en["enr_confidence"] == "Review (name only)").sum())
    fe = int(((en["enr_status"] == "Enriched") & (en["enr_source"] == "FullEnrich")).sum())
    other = int((en["enr_status"] == "Enriched").sum()) - fe
    rows = [
        ("Report", "iFax — Active Customers (All Enrichment Sources, Re-bucketed)"),
        ("Total active accounts", f"{n:,}"),
        ("Enrichment sources", "FullEnrich + Other_Paying + High_LTV + Activated_Freemiums"),
        ("Matched — 100% (exact domain)", f"{ex:,} ({ex/n*100:.1f}%)"),
        ("Matched — name (review, FullEnrich only)", f"{rv:,} ({rv/n*100:.1f}%)"),
        ("Total enriched", f"{ex+rv:,} ({(ex+rv)/n*100:.1f}%)"),
        ("  · via FullEnrich", f"{fe:,}"),
        ("  · via past-enrichment files (new)", f"{other:,}"),
        ("Accounts re-bucketed (size corrected)", f"{moved:,}"),
        ("Size coverage BEFORE enrichment", f"{size_b:,} ({size_b/n*100:.1f}%)"),
        ("Size coverage AFTER enrichment", f"{size_a:,} ({size_a/n*100:.1f}%)"),
        ("", ""),
        ("NOTES", ""),
        ("Source tag", "Every enriched row carries enr_source. Domains found in FullEnrich keep that "
                       "(richest) record; new domains come from the past-enrichment files (name-only mostly)."),
        ("Confidence", "Only exact domain matches are 100% and eligible for re-bucketing. Name-only "
                       "matches (FullEnrich) are flagged 'review'. No fuzzy guesses."),
        ("Data fix", "Corrupted 'headcount range' (1-10 shown as 2026-02-10; '10001' without +) decoded."),
    ]
    return pd.DataFrame(rows, columns=["Metric", "Value"])


def build(csv_dir, fullenrich_path, extra_files, out_path):
    df = base.load(csv_dir)
    dom_best, name_unique, dom_band, dom_source, stats = build_multi_index(fullenrich_path, extra_files)
    en = enrich(df, dom_best, name_unique, dom_band, dom_source)

    df_final = en.copy()
    df_final["size_bucket"] = en["size_bucket_final"]
    size_b = int((en["original_bucket"] != "Not captured").sum())
    size_a = int((en["size_bucket_final"] != "Not captured").sum())
    moved = int((en["moved"] == "Yes").sum())

    sheets = [
        ("Overview", overview(en, moved, size_b, size_a, dom_source)),
        ("Coverage Before vs After", rb.coverage_before_after(en)),
        ("Enrichment Coverage", rb.enrichment_coverage(en)),
        ("Coverage by Source", coverage_by_source(en)),
        ("Re-bucketing Log", rb.rebucket_log(en)),
        ("Segment Summary", base.t_segment_summary(df_final)),
        ("Coverage - Overall", base.t_coverage_overall(df)),
        ("Coverage - by Segment", base.t_coverage_by_segment(df_final)),
        ("Missing-Data & Gaps", base.t_gaps(df_final)),
        ("Completeness Score", base.t_completeness(df_final)),
        ("Plan & Tier Breakdown", base.t_tier(df_final)),
        ("Top Accounts", base.t_top_accounts(df_final)),
    ]
    base_cols = ["customer_id", "best_company", "company", "domain", "email",
                 "first_name", "last_name", "job_title", "state", "country", "zip",
                 "active_mrr", "active_arr", "plans", "tier", "term", "is_trial",
                 "currency", "csm", "origin", "partner", "completeness_0_7"]
    for seg in base.SIZE_ORDER:
        s = en[en["size_bucket_final"] == seg]
        if s.empty:
            continue
        s = s.sort_values("active_mrr", ascending=False)
        tab = s[base_cols + FINAL_ENRICH_COLS + META_COLS].rename(
            columns={"active_mrr": "MRR", "active_arr": "ARR"}).reset_index(drop=True)
        sheets.append(("Not Captured" if seg == "Not captured" else seg, tab))

    _write(sheets, out_path, green=set(FINAL_ENRICH_COLS), amber=set(META_COLS))
    _verify(en, size_b, size_a, moved, dom_source)
    return en


def _write(sheets, out_path, green, amber):
    ENR_HDR = PatternFill("solid", fgColor="375623"); ENR_BODY = PatternFill("solid", fgColor="E2EFDA")
    META_HDR = PatternFill("solid", fgColor="806000"); META_BODY = PatternFill("solid", fgColor="FFF2CC")
    BLUE = PatternFill("solid", fgColor="1F3864")
    with pd.ExcelWriter(out_path, engine="openpyxl") as writer:
        for name, d in sheets:
            d.to_excel(writer, sheet_name=name[:31], index=False)
        wb = writer.book
        for name, d in sheets:
            ws = wb[name[:31]]
            for c in range(1, d.shape[1] + 1):
                h = str(d.columns[c - 1]); cell = ws.cell(row=1, column=c)
                cell.font = Font(bold=True, color="FFFFFF")
                cell.alignment = Alignment(horizontal="left", vertical="center")
                g, a = h in green, h in amber
                cell.fill = ENR_HDR if g else META_HDR if a else BLUE
                ws.column_dimensions[get_column_letter(c)].width = min(
                    max([len(h)] + [len(str(v)) for v in d.iloc[:, c - 1].tolist()[:400]]) + 2, 55)
                fmt = '0.0"%"' if "%" in h else ('$#,##0.00' if base._money_col(h) else None)
                for r in range(2, d.shape[0] + 2):
                    cc = ws.cell(row=r, column=c)
                    if fmt:
                        cc.number_format = fmt
                    if g:
                        cc.fill = ENR_BODY
                    elif a:
                        cc.fill = META_BODY
            ws.freeze_panes = "A2"
            if d.shape[1] and d.iloc[:, 0].astype(str).eq("TOTAL").any():
                i = d.index[d.iloc[:, 0].astype(str) == "TOTAL"][0]
                for c in range(1, d.shape[1] + 1):
                    ws.cell(row=i + 2, column=c).font = Font(bold=True)


def _verify(en, size_b, size_a, moved, dom_source):
    n = len(en)
    assert n == 7431 and en["customer_id"].nunique() == n
    assert int(en["size_bucket_final"].value_counts().sum()) == n
    mv = en[en["moved"] == "Yes"]
    assert (mv["enr_confidence"] == "100%").all()
    ex = int((en["enr_confidence"] == "100%").sum())
    fe = int(((en["enr_status"] == "Enriched") & (en["enr_source"] == "FullEnrich")).sum())
    new = int((en["enr_status"] == "Enriched").sum()) - fe - int((en["enr_confidence"] == "Review (name only)").sum())
    print(f"  verify OK: {n} accounts, no row loss")
    print(f"  matched 100%: {ex} ({ex/n*100:.1f}%) | re-bucketed: {moved}")
    print(f"  by source: {en[en['enr_status']=='Enriched']['enr_source'].value_counts().to_dict()}")
    print(f"  size coverage before {size_b} ({size_b/n*100:.1f}%) -> after {size_a} ({size_a/n*100:.1f}%)")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--dir", required=True)
    ap.add_argument("--fullenrich", required=True)
    ap.add_argument("--extra", nargs="+", required=True, help="label=path pairs")
    ap.add_argument("--out", default="icp_overview_all_sources.xlsx")
    a = ap.parse_args()
    extra = [tuple(x.split("=", 1)) for x in a.extra]
    build(a.dir, a.fullenrich, extra, a.out)
    print(f"  wrote {a.out}")


if __name__ == "__main__":
    main()
