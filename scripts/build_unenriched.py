"""Build the 'left to enrich' workbook: every active customer that NO FullEnrich
source (main file + the 3 past exports) could match — bucketed by company size.

Original Chargebee fields only; one tab per size bucket + a summary.
"""
import sys
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import build_icp_sheet as base
import build_final_workbook as bf

ORIG_FIELDS = ["customer_id", "company", "company_from_domain", "domain", "domain_from_email",
               "email", "first_name", "last_name", "job_title", "state", "country", "zip",
               "active_mrr", "active_arr", "plans", "currency", "csm", "origin", "partner"]
HDR = PatternFill("solid", fgColor="1F3864")


def build(csv_dir, fullenrich, extra_files, out_path):
    df = base.load(csv_dir)
    dom_best, name_unique, dom_band, dom_source, _ = bf.build_multi_index(fullenrich, extra_files)
    en = bf.enrich(df, dom_best, name_unique, dom_band, dom_source)

    un = en[en["enr_status"] == "Not enriched"].copy()
    total = len(en)

    # summary per original bucket
    rows = []
    for seg in base.SIZE_ORDER:
        segdf = en[en["size_bucket"] == seg]
        u = un[un["size_bucket"] == seg]
        rows.append({
            "Company size bucket": seg,
            "Total active": len(segdf),
            "Enriched": len(segdf) - len(u),
            "NOT enriched (left to do)": len(u),
            "% left": round(len(u) / len(segdf) * 100, 1) if len(segdf) else 0.0,
            "MRR left to enrich": round(u["active_mrr"].sum(), 2),
        })
    summ = pd.DataFrame(rows)
    summ.loc[len(summ)] = {
        "Company size bucket": "TOTAL", "Total active": total,
        "Enriched": total - len(un), "NOT enriched (left to do)": len(un),
        "% left": round(len(un) / total * 100, 1),
        "MRR left to enrich": round(un["active_mrr"].sum(), 2),
    }

    sheets = [("Summary — Left to Enrich", summ)]
    for seg in base.SIZE_ORDER:
        s = un[un["size_bucket"] == seg]
        if s.empty:
            continue
        tab = s[ORIG_FIELDS].sort_values("active_mrr", ascending=False).reset_index(drop=True)
        sheets.append(("Not Captured" if seg == "Not captured" else seg, tab))

    with pd.ExcelWriter(out_path, engine="openpyxl") as w:
        for name, d in sheets:
            d.to_excel(w, sheet_name=name[:31], index=False)
        wb = w.book
        for name, d in sheets:
            ws = wb[name[:31]]
            for c in range(1, d.shape[1] + 1):
                cell = ws.cell(row=1, column=c)
                cell.fill = HDR
                cell.font = Font(bold=True, color="FFFFFF")
                cell.alignment = Alignment(horizontal="left", vertical="center")
                L = get_column_letter(c)
                vals = [len(str(v)) for v in d.iloc[:, c - 1].tolist()[:400]]
                ws.column_dimensions[L].width = min(max([len(str(d.columns[c - 1]))] + vals) + 2, 45)
                fmt = '0.0"%"' if "%" in str(d.columns[c - 1]) else (
                    '$#,##0.00' if base._money_col(str(d.columns[c - 1])) else None)
                if fmt:
                    for r in range(2, d.shape[0] + 2):
                        ws.cell(row=r, column=c).number_format = fmt
            ws.freeze_panes = "A2"
            if d.iloc[:, 0].astype(str).eq("TOTAL").any():
                i = d.index[d.iloc[:, 0].astype(str) == "TOTAL"][0]
                for c in range(1, d.shape[1] + 1):
                    ws.cell(row=i + 2, column=c).font = Font(bold=True)

    # verify
    assert len(un) + (en["enr_status"] == "Enriched").sum() == total
    print(f"wrote {out_path}")
    print(f"  total active {total} | NOT enriched (left) {len(un)} ({len(un)/total*100:.1f}%)")
    print(summ[["Company size bucket", "NOT enriched (left to do)"]].to_string(index=False))


if __name__ == "__main__":
    D = "/root/.claude/uploads/815bd82b-022c-5eff-8466-360a03104051/"
    build(D, D + "15cd01b6-Chargebee_Company_Enrichment_.xlsx",
          [("Other_Paying", D + "69ada757-Copy_of_Other_Paying.xlsx"),
           ("High_LTV", D + "b3d7526a-Copy_of_High_LTV.xlsx"),
           ("Activated_Freemiums", D + "09a09842-Copy_of_Activated_Freemiums.xlsx")],
          sys.argv[1] if len(sys.argv) > 1 else "Left_To_Enrich.xlsx")
