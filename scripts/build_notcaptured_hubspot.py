"""Build the Not-Captured -> HubSpot company-size lookup workbook."""
import json
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

SP = "/tmp/claude-0/-home-user-Enrichment/815bd82b-022c-5eff-8466-360a03104051/scratchpad"
ORDER = ["1-10", "11-50", "51-200", "201-500", "501-1000", "1001-5000", "5001-10000", "10001+"]
HDR = PatternFill("solid", fgColor="1F3864")


def load():
    acc = pd.read_csv(f"{SP}/not_captured_hubspot_lookup.csv", dtype=str)
    acc["MRR"] = pd.to_numeric(acc["active_mrr"], errors="coerce").fillna(0).round(2)
    return acc


def t_summary(acc):
    n = len(acc)
    sized = int((acc["hs_status"] == "HubSpot size found").sum())
    nosize = int((acc["hs_status"] == "company exists, no size").sum())
    nocomp = int((acc["hs_status"] == "no HubSpot company").sum())
    nodom = int((acc["hs_status"] == "(no company domain)").sum())
    rows = [
        ["Not-captured accounts looked up", n, ""],
        ["HubSpot HAS a company size", sized, f"{sized/n*100:.1f}%"],
        ["HubSpot company exists but size field blank", nosize, f"{nosize/n*100:.1f}%"],
        ["Real domain but NO HubSpot company", nocomp, f"{nocomp/n*100:.1f}%"],
        ["Free-email only (gmail/etc.) — no company domain", nodom, f"{nodom/n*100:.1f}%"],
        ["", "", ""],
        ["Flagged for manual research (no usable size)", n - sized, f"{(n-sized)/n*100:.1f}%"],
        ["  · researchable (has a real company domain)", nocomp + nosize, ""],
        ["  · gmail/free-email only (hardest)", nodom, ""],
        ["", "", ""],
        ["NOTE — SLG / process coverage", "", "Only ~11% of not-captured accounts have a HubSpot size; most "
         "came from HubSpot Insights auto-enrichment, not manual sales entry."],
        ["NOTE — data integrity", "", "HubSpot's custom 'company_size' field frequently contradicts the numeric "
         "'numberofemployees' (e.g. 533 employees tagged '11-50'); bucketing uses numberofemployees."],
        ["NOTE — concentration", "", "37 of the 101 sized accounts share one domain (nationalbreathefree.com, "
         "a multi-location group), inflating the 51-200 bucket."],
    ]
    return pd.DataFrame(rows, columns=["Metric", "Count", "Note"])


def t_distribution(acc):
    sized = acc[acc["hs_status"] == "HubSpot size found"]
    d = sized["hs_bucket"].value_counts().reindex(ORDER).fillna(0).astype(int)
    mrr = sized.groupby("hs_bucket")["MRR"].sum().reindex(ORDER).fillna(0).round(2)
    out = pd.DataFrame({"Org-size bucket (HubSpot)": ORDER, "Accounts": d.values, "MRR": mrr.values})
    out.loc[len(out)] = ["TOTAL", int(d.sum()), round(float(mrr.sum()), 2)]
    return out


def t_sized(acc):
    s = acc[acc["hs_status"] == "HubSpot size found"].copy()
    s = s.sort_values(["hs_bucket", "MRR"], ascending=[True, False])
    cols = ["customer_id", "company", "email", "lookup_domain", "hs_company", "hs_employees", "hs_bucket", "MRR", "csm", "origin"]
    return s[cols].rename(columns={"lookup_domain": "domain", "hs_company": "hubspot_company",
                                   "hs_employees": "hubspot_employees", "hs_bucket": "hubspot_size_bucket"})


def t_flag(acc):
    f = acc[acc["hs_status"] != "HubSpot size found"].copy()
    f["research_path"] = f["lookup_domain"].fillna("").map(
        lambda d: "domain lookup" if d else "no company domain (gmail) — hardest")
    f = f.sort_values(["research_path", "MRR"], ascending=[True, False])
    cols = ["customer_id", "company", "email", "lookup_domain", "hs_status", "research_path", "MRR", "csm", "origin"]
    return f[cols].rename(columns={"lookup_domain": "domain", "hs_status": "hubspot_status"})


def write(path):
    acc = load()
    sheets = [
        ("Summary", t_summary(acc)),
        ("Distribution", t_distribution(acc)),
        ("HubSpot Size Found", t_sized(acc)),
        ("Flag - Manual Research", t_flag(acc)),
    ]
    with pd.ExcelWriter(path, engine="openpyxl") as w:
        for name, d in sheets:
            d.to_excel(w, sheet_name=name[:31], index=False)
        wb = w.book
        for name, d in sheets:
            ws = wb[name[:31]]
            for c in range(1, d.shape[1] + 1):
                cell = ws.cell(row=1, column=c)
                cell.fill = HDR
                cell.font = Font(bold=True, color="FFFFFF")
                cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
                L = get_column_letter(c)
                vals = [len(str(v)) for v in d.iloc[:, c - 1].tolist()[:400]]
                ws.column_dimensions[L].width = min(max([len(str(d.columns[c - 1]))] + vals) + 2, 60)
            ws.freeze_panes = "A2"
            if d.iloc[:, 0].astype(str).eq("TOTAL").any():
                i = d.index[d.iloc[:, 0].astype(str) == "TOTAL"][0]
                for c in range(1, d.shape[1] + 1):
                    ws.cell(row=i + 2, column=c).font = Font(bold=True)
    print("wrote", path)


if __name__ == "__main__":
    import sys
    write(sys.argv[1] if len(sys.argv) > 1 else "NotCaptured_HubSpot_Lookup.xlsx")
